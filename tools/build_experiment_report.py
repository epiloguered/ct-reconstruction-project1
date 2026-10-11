"""生成同源Word与Markdown课程报告；图表来自build_report_assets.py的已保存结果汇总。"""
from pathlib import Path
import csv, json, hashlib, re
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DEST=ROOT/'report'; FINAL=ROOT/'outputs/final'; BUILD=ROOT/'outputs/report_build'
TITLE='CT图像重建与稀疏角度优化实验报告'
FILE_STEM=TITLE+'修订版'
BLOCKS=[]; TABLE_DATA=[]


def read(name):return list(csv.DictReader((FINAL/name).open(encoding='utf-8-sig',newline='')))
def add(kind,**kw):BLOCKS.append({'kind':kind,**kw})
def p(text):add('p',text=text)
def bullet(text):add('bullet',text=text)
def h(text,level=1):add('h',text=text,level=level)
def page():add('page')
def eq(text,parts=None):add('eq',text=text,parts=parts or [text])
def table(caption,headers,rows,widths=None):
    add('table',caption=caption,headers=headers,rows=rows,widths=widths)
    TABLE_DATA.append({'caption':caption,'headers':headers,'rows':rows})
def fig(path,caption,width=16):add('fig',path=path,caption=caption,width=width)
def n(v):return f'{float(v):.4f}'
def ss(v):return f'{float(v):.2e}' if abs(float(v))<.00005 else f'{float(v):.4f}'
def triple(r):return [n(r['nrmse']),f"{float(r['psnr_db']):.2f}",ss(r['ssim'])]
LABEL={'LSQR':'LSQR','ART_fixed':'ART 第5轮','LS_L2':'LS+L2','LS_L1':'LS+L1','LS_TV':'LS+TV',
       'ART_L2':'ART+L2','ART_L1':'ART+L1','ART_TV':'ART+TV','tv':'非负TV','wave':'小波+TV','pnp':'PnP-LC','pnp_np':'PnP常规'}
SHORT={'SheppLogan64':'SL64','XCAT128':'XCAT128'}


def compose():
    core=read('baseline_regularization.csv');classic=read('sparse_classical.csv');fixed=read('non_dip_fixed.csv')
    core_get=lambda name,snr,m:next(r for r in core if r['dataset']==name and int(r['snr_label_db'])==snr and r['method']==m)
    github='https://github.com/epiloguered/ct-reconstruction-project1'
    h(TITLE,0)
    p('课程 最优化理论与算法    实验 Project 1')
    p('姓名 __________________    学号 __________________')
    p('日期 2026年10月11日')
    p('本报告总结CT图像的基础重建、正则化与稀疏角度实验，比较各方法的像素误差和结构保持能力。')
    add('link',label='项目代码仓库',url=github)
    p('代码、环境配置、实验参数和完整结果均在该仓库中整理，正文聚焦方法、结果与结论。')
    add('toc')
    page()
    h('摘要',2)
    p('本实验研究带噪声线性投影模型下的CT图像重建。使用课程给定的SheppLogan64与XCAT128数据，在15、20、25、30 dB四档信噪比下完成最小二乘和ART基线，以及L2、L1、TV正则化；进一步减少投影角度，比较非负TV、TGV、小波＋TV、DIP及PnP-BM3D。结果表明，无正则化最小二乘容易出现图像误差放大；TV在LS与ART两个系列的全部八个数据条件中均取得最低NRMSE。PnP-BM3D改善了部分少角度重建，但10角高SSIM案例仍比原90角LS+TV具有更高像素误差。报告同时汇总每种方法的最佳固定终点结果，展示角度数、噪声水平与重建质量的关系。')
    p('关键词 CT重建；最小二乘；ART；总变差；稀疏角度；PnP-BM3D')
    h('一 实验目的与数据')
    bullet('基础任务：使用给定投影矩阵A与带噪声观测b，通过LS和ART恢复图像，并比较L2、L1、TV正则化的效果。')
    bullet('扩展任务：减少投影方向，研究欠定系统下的重建质量，并探索尽量保持质量的少角度方案。质量参照采用原90角LS+TV和ART+TV。')
    table('表1 实验数据维度',['数据集','图像大小','原A维度','每方向射线','10角A维度'],[
        ['SheppLogan64','64×64','8190×4096','91','910×4096'],
        ['XCAT128','128×128','16290×16384','181','1810×16384']], [3.1,2.3,3.7,2.8,4.1])
    p('数据由课程提供。下文SL64为SheppLogan64的简写，GT表示真实图像。原90个方向为0、2、…、178度；图像按列优先F顺序展开为向量x。每条射线只经过部分像素，因此A中绝大多数元素为0。采用CSR稀疏矩阵只存储非零元素及其位置，能够节省内存，并加快迭代中的投影Ax与残差反投影Aᵀ(Ax−b)。')
    eq('b = A x_gt + ε；SNR = 20 log₁₀(‖A x_gt‖₂ / ‖ε‖₂)')
    p('15–30 dB为输入信噪比：相同信号下，SNR越高，相对噪声越小。四档观测均使用给定文件。')
    page()
    h('二 基础重建与正则化方法')
    h('2.1 最小二乘与LSQR',2)
    eq('x_LS ∈ arg min_x ½ ‖Ax − b‖₂²')
    p('最小二乘LS通过最小化投影误差求解图像；LSQR是求解大型稀疏最小二乘问题的数值算法，依据Paige与Saunders的工作[1]。本实验调用SciPy LSQR，atol=btol=10⁻⁵、迭代上限10000、条件数上限10⁸，采用零初始化。')
    h('2.2 代数重建技术ART',2)
    p('ART依次利用每条射线纠正图像，依据Gordon等的代数重建技术[2]。下式aᵢ为A的第i行，松弛参数ω=1；每行更新后将负像素截为0，零范数行跳过更新。实验记录40轮，主结果采用固定第5轮。')
    eq('x ← max{0, x + ω (bᵢ − aᵢx) aᵢᵀ / ‖aᵢ‖₂²}')
    h('2.3 L2 L1与TV正则化',2)
    eq('min_x ½ ‖Ax − b‖₂² + λ R(x)')
    eq('R_L2(x)=½‖x‖₂²；R_L1(x)=‖x‖₁；TV(x)=Σᵢⱼ √((Dₕx)ᵢⱼ²+(Dᵥx)ᵢⱼ²)')
    bullet('L2惩罚较大的像素能量；LS+L2使用damp=√λ的LSQR。')
    bullet('像素L1鼓励像素值稀疏；LS+L1使用FISTA[4]与软阈值更新。')
    bullet('TV惩罚相邻像素梯度，其理论依据为Rudin、Osher与Fatemi的总变差模型[3]。LS+TV使用FISTA[4]，TV近端采用Chambolle算法[16]。')
    bullet('ART+L2／L1／TV在每轮ART投影后加入相应收缩或TV处理，并与原ART的固定5轮结果比较。')
    table('表2 原90角正则化方案的迭代设置',['方案','迭代设置','图像约束'],[
        ['LS+L2','LSQR最多5000次','无非负约束'],
        ['LS+L1','FISTA 1200步','无非负约束'],
        ['LS+TV','FISTA 600步；TV近端最多80步','无非负约束'],
        ['ART+L2/L1/TV','固定5轮；每轮后加入先验处理','非负约束']], [3.5,8.0,4.5])
    page()
    h('三 少角度扩展方法')
    h('3.1 稀疏角度采样',2)
    p('少角度实验同步抽取A和b的测量行，保持像素数不变。对K个方向，原方向编号取floor(90j/K)，j=0,…,K−1；每方向保留全部探测器。例如10角为0、18、…、162度。稀疏角度CT使用TV的思路参考Sidky等[7]。')
    bullet('传统扩展覆盖90、45、30、15角，以及15、20、25、30 dB全部四档SNR。')
    bullet('DIP与后续改进选取20、30 dB；非负TV、小波和PnP主要测试30、15、10、6、3角，SL64／30 dB另补测7、8、9角。')
    h('3.2 非负TV TGV与小波先验',2)
    p('非负TV在TV重建中加入x≥0约束。二阶TGV参考Bredies等[11]，同时约束一阶变化与二阶变化；小波＋TV联合利用小波系数和图像梯度的稀疏性，小波阈值理论参考Donoho[12]。本实验采用db2小波，优化使用Condat型原始—对偶分裂[13]。')
    h('3.3 PnP与BM3D',2)
    p('Plug-and-Play（PnP）在迭代中交替执行投影拟合与图像去噪，参考Venkatakrishnan等[8]。本实验使用ADMM型更新，将BM3D作为去噪器；BM3D利用相似图像块的协同滤波，参考Dabov等[9]。实际库为bm3d 4.0.3[15]，比较LC配置与常规配置。')
    eq('xᵏ⁺¹ = arg min_x ½‖Bx−d‖₂² + ½ρ‖x−zᵏ+uᵏ‖₂²')
    eq('zᵏ⁺¹ = max{D_σ(xᵏ⁺¹+uᵏ),0}；uᵏ⁺¹ = uᵏ+xᵏ⁺¹−zᵏ⁺¹')
    p('变量按GT极差s归一化，B=A/√L、d=b/(s√L)，输出s·z；Dσ为BM3D。投影子问题用预条件CG近似求解。PnP以同条件12000步非负TV结果为起点，9组参数各运行40步选参，再续至160步。')
    h('3.4 深度图像先验DIP',2)
    p('DIP依据Ulyanov等[10]，将图像写为未训练网络的输出x=s·fθ(z)，针对单份投影优化网络参数。本实验采用四尺度卷积网络、Softplus非负输出、固定随机输入、Adam优化和2000步训练；种子为0、1、2。DIP+TV在拟合损失中加入TV项，完整结果见附录。')
    page()
    h('四 评价指标与实验协议')
    h('4.1 指标选择',2)
    p('采用NRMSE、PSNR和SSIM共同评价重建图像。平方误差和PSNR的意义参考Wang与Bovik[5]，SSIM参考Wang等[6]；实际指标实现参考scikit-image文档[14]。')
    table('表3 评价指标及方向',['指标','评价内容','方向'],[
        ['NRMSE','相对像素误差：误差L2范数除以GT的L2范数','越小越好'],
        ['PSNR','由均方误差计算的峰值信噪比，单位dB','越大越好'],
        ['SSIM','比较图像局部亮度、对比度与结构相似性','越大越好']], [2.5,10.5,3.0])
    p('NRMSE采用GT范数归一化；PSNR的强度范围取GT最大值减最小值。SSIM采用scikit-image默认7×7均匀窗口。另记录相对投影残差‖Ax−b‖₂/‖b‖₂，用于观察投影拟合程度。所有方法使用相同的指标实现和未经显示裁剪的保存数组。')
    h('4.2 结果比较方式',2)
    bullet('正则化和传统扩展在给定网格中按GT NRMSE选参数，PSNR与SSIM来自同一组参数。')
    bullet('DIP主结果为固定2000步、三个种子的均值与样本标准差；非负TV及小波主结果为12000步，PnP主结果为160步。')
    bullet('主表保留各方法实际的约束、初始化及迭代设置；PnP计算还包括非负TV起点。固定终点结果与根据GT另选保存终点的结果分别报告。')
    bullet('少角度质量分别对照原90角LS+TV和ART+TV。传统扩展的90角重跑只用于该阶段内部的角度趋势比较。')
    page()
    h('五 无正则化基线结果')
    rows=[]
    for name in ('SheppLogan64','XCAT128'):
        for snr in (15,20,25,30):
            for method in ('LSQR','ART_fixed'):
                r=core_get(name,snr,method)
                rows.append([SHORT[name],snr,LABEL[method],*triple(r),n(r['relative_residual'])])
    table('表4 原90角无正则化基线结果',['数据集','SNR','方法','NRMSE↓','PSNR↑ dB','SSIM↑','投影残差↓'],rows,[2.2,1.2,2.7,2.5,2.6,2.4,2.4])
    bullet('ART固定第5轮在两个数据集、全部四档SNR下均比无正则化LSQR具有更低NRMSE。30 dB时，SL64的两者分别为0.1340与0.5926，XCAT128分别为0.1760与18.1019。')
    bullet('较低投影残差并不对应较低图像误差。XCAT128／15 dB的LSQR投影残差仅0.0677，但NRMSE达到121.7020、PSNR为−31.43 dB。')
    bullet('SNR从15提高至30 dB时，两种基线的NRMSE整体下降、PSNR和SSIM整体上升，说明相对噪声减小后重建质量改善。')
    page()
    h('六 正则化实验结果')
    methods=('LSQR','LS_L2','LS_L1','LS_TV','ART_fixed','ART_L2','ART_L1','ART_TV')
    for idx,name in enumerate(('SheppLogan64','XCAT128')):
        if idx:page()
        h(f'6.{idx+1} {name}结果',2)
        rows=[[snr,LABEL[m],*triple(core_get(name,snr,m))] for snr in (15,20,25,30) for m in methods]
        table(f'表{5+idx} {name}原90角基线与正则化结果',['SNR dB','方法','NRMSE↓','PSNR↑ dB','SSIM↑'],rows,[2.0,4.4,3.2,3.2,3.2])
        if idx==0:
            bullet('四档SNR下，LS+TV与ART+TV分别取得各自算法系列最低NRMSE；30 dB时两者为0.0688与0.0943。')
            bullet('20 dB时LS+L1的SSIM为0.7135，略高于LS+TV的0.7077；TV仍具有更低像素误差。')
        else:
            bullet('TV同样在四档SNR下取得各自算法系列最低NRMSE；30 dB时LS+TV由原LSQR的18.1019降至0.1066。')
            bullet('15 dB时LS+L1的SSIM为0.6819，高于TV的0.5219，但其NRMSE为0.8665，高于TV的0.2269。')
    page()
    h('6.3 正则化作用分析',2)
    fig('figures/regularization_images.png','图1 30 dB时LS系列的正则化对照。每个数据集共用GT灰度范围。',16)
    bullet('L2明显抑制无正则化LS的大幅度误差，但图像仍有噪声。')
    bullet('像素L1与TV使用不同先验：前者约束像素值稀疏性，后者约束梯度稀疏性。给定图像中TV对平坦区域和边缘的恢复更好。')
    bullet('两个数据集、四档SNR共八个条件中，TV在LS与ART系列均取得最低NRMSE；SSIM的最佳方法随条件变化。')
    page()
    h('七 稀疏角度传统方法结果')
    rows=[]
    for name in ('SheppLogan64','XCAT128'):
        for snr in (15,20,25,30):
            rr=[next(r for r in classic if r['dataset']==name and int(r['snr_label_db'])==snr and int(r['angle_count'])==k and r['method']=='LS_TV') for k in (90,45,30,15)]
            rows.append([SHORT[name],snr,*[n(r['nrmse']) for r in rr]])
    table('表7 传统扩展LS+TV的NRMSE',['数据集','SNR dB','90角重跑','45角','30角','15角'],rows,[2.4,1.6,3.0,3.0,3.0,3.0])
    rows=[]
    for name in ('SheppLogan64','XCAT128'):
        for k in (90,45,30,15):
            for m in ('LS_TV','ART_TV'):
                r=next(r for r in classic if r['dataset']==name and int(r['snr_label_db'])==30 and int(r['angle_count'])==k and r['method']==m)
                rows.append([SHORT[name],k,LABEL[m],*triple(r)])
    table('表8 30 dB传统稀疏角度TV方案的三指标',['数据集','角度','方法','NRMSE↓','PSNR↑ dB','SSIM↑'],rows,[2.4,1.3,3.0,3.1,3.1,3.1])
    bullet('本阶段覆盖全部四档SNR。90角行使用该阶段重跑结果，与45、30、15角形成同一协议下的对照。')
    bullet('LS+TV的NRMSE在各SNR下均随90→45→30→15角减少而上升。30 dB时，SL64由0.0688升至0.2912，XCAT128由0.1066升至0.1885。')
    bullet('SSIM不严格随角度减少而下降。例如XCAT128／30 dB的45角LS+TV SSIM为0.8874，高于90角的0.8237，但NRMSE仍更高。')
    page()
    h('八 非负TV 小波与PnP重建改进')
    p('本节比较非负TV与PnP-BM3D。非负TV和小波＋TV取12000步终点，PnP取160步终点；参数按GT NRMSE选择。小波及3／6角结果同时展示在质量曲线中。')
    rows=[]
    for name in ('SheppLogan64','XCAT128'):
        for snr in (20,30):
            for k in (30,15,10):
                for m in ('tv','pnp'):
                    r=next(r for r in fixed if r['dataset']==name and int(r['snr_label_db'])==snr and int(r['angle_count'])==k and r['method']==m)
                    rows.append([SHORT[name],snr,k,LABEL[m],*triple(r)])
    table('表9 非负TV与PnP-BM3D LC固定终点对照',['数据集','SNR','角度','方法','NRMSE↓','PSNR↑ dB','SSIM↑'],rows,[2.2,1.1,1.1,3.2,2.8,2.8,2.8])
    bullet('SL64的六个表列条件下，PnP-LC均比非负TV具有更低NRMSE。30 dB／15角时，NRMSE由0.2048降至0.1336，SSIM由0.8540升至0.9380。')
    bullet('XCAT128的结果取决于噪声与角度。30 dB／15角时，PnP-LC NRMSE为0.1948，高于非负TV的0.1771，但SSIM略高。')
    page()
    h('8.1 角度与质量的权衡',2)
    fig('figures/quality_angles_fixed.png','图2 各方法固定终点NRMSE与角度数。虚线为原90角LS+TV／ART+TV参照；常规BM3D配置仅在部分30 dB条件补测。',16)
    p('进一步在全部已保存候选和续迭代终点中，寻找同时满足“NRMSE不高于90角参照、SSIM不低于同一参照”的最少已测角度。该表额外使用GT选择终点，与表9固定终点的口径不同。')
    audit=json.loads((ROOT/'outputs/min_angles/fullview_audit/audit.json').read_text(encoding='utf-8'))
    rows=[]
    for entry in audit['thresholds']:
        r=entry['reference'];b=entry['least_tested_match']
        rows.append([SHORT[r['dataset']],r['snr'],LABEL[r['method']],str(b['angle_count']) if b else '未达到',
                     (LABEL[b['method']]+' '+str(b['iterations'])+'步') if b else '—',n(b['nrmse']) if b else '—',ss(b['ssim']) if b else '—'])
    table('表10 同时达到原90角NRMSE和SSIM的最少已测角度',['数据集','SNR','90角参照','最少角度','方案','NRMSE','SSIM'],rows,[2.2,1.1,2.6,2.1,2.7,2.7,2.6])
    bullet('相对原90角ART+TV，SL64的20与30 dB可在30角达标，XCAT128／20 dB可在10角达标。')
    bullet('相对原90角LS+TV，四个20／30 dB条件均未在本次3–30角保存结果中同时达标。')
    page()
    h('8.2 高SSIM案例分析',2)
    rows=[['90角 '+LABEL[m],*triple(core_get('SheppLogan64',30,m))] for m in ('LS_TV','ART_TV')]
    npcase=next(r for r in fixed if r['dataset']=='SheppLogan64' and int(r['snr_label_db'])==30 and int(r['angle_count'])==10 and r['method']=='pnp_np')
    rows.append(['10角 PnP常规160步',*triple(npcase)])
    table('表11 SL64的30 dB常规BM3D配置补测',['方案','NRMSE↓','PSNR↑ dB','SSIM↑'],rows,[6.4,3.2,3.2,3.2])
    p('10角常规配置沿用LC选中的ρ=0.03、σ=0.03。保存结果NRMSE为0.114672、SSIM为0.944047。')
    fig('figures/ssim_error_case.png','图3 SL64／30 dB的重建与误差。下排绝对误差以GT极差归一化，色阶固定0–0.3。',16)
    region=read('region_diagnostics.csv');rows=[]
    for m in ('LS_TV','ART_TV','TV10','PnP10'):
        r=next(r for r in region if r['dataset']=='SheppLogan64' and r['method']==m)
        label={'LS_TV':'90角LS+TV','ART_TV':'90角ART+TV','TV10':'10角非负TV','PnP10':'10角常规PnP'}[m]
        rows.append([label,n(r['background_ssim_map_mean']),n(r['foreground_ssim_map_mean']),n(r['edge_rmse'])])
    table('表12 同一案例的区域诊断',['方案','背景SSIM↑','前景SSIM↑','边缘RMSE↓'],rows,[6.4,3.2,3.2,3.2])
    bullet('10角PnP相比90角LS+TV，NRMSE高约66.7%，PSNR低约4.44 dB，SSIM略高。')
    bullet('零背景占图像58.84%。PnP背景SSIM更高，但前景SSIM更低，边缘RMSE约为90角LS+TV的2.06倍；高全图SSIM主要伴随背景改善，小结构和边缘仍有损失。')
    page()
    h('九 各方法最佳结果主表')
    p('每个数据集、每种方法取已测条件中的最低NRMSE，三指标同源，选中条件均为30 dB。基础固定90角，扩展少于90角；各方法角度覆盖不同。DIP取三种子均值，其余按规定终点。本表用于成果汇总，同条件比较见表9，最少角度达标见表10。')
    best=read('best_by_method.csv')
    rows=[['基础' if r['stage']=='basic' else '扩展',SHORT[r['dataset']],r['method_label'],r['angle_count'],*triple(r)] for r in best]
    table('表13 基础与扩展各方法的最佳固定终点结果 SNR均为30 dB',['部分','数据集','方法','角度','NRMSE↓','PSNR↑ dB','SSIM↑'],rows,[1.0,2.2,4.1,1.0,2.6,2.6,2.5])
    page()
    h('十 讨论与结论')
    h('10.1 基础实验',2)
    bullet('无正则化LS容易放大带噪声观测中的误差；采用非负约束、固定5轮的ART在八个条件下均比LSQR具有更低NRMSE。')
    bullet('L2、L1和TV整体改善了LS重建，其中TV在八个数据条件中均取得LS与ART各系列最低NRMSE。')
    bullet('PSNR、NRMSE与SSIM反映不同侧面。L1在部分条件有较高SSIM，但像素误差仍高于TV；投影残差低也不保证图像误差低。')
    h('10.2 扩展实验',2)
    bullet('减少投影方向会增加重建难度。传统LS+TV在四档SNR下均随角度减少出现NRMSE上升。')
    bullet('非负TV、小波＋TV与PnP-BM3D改善了部分少角度结果。PnP对SL64改善更明显，XCAT128部分条件仍由非负TV更好。')
    bullet('10角PnP可以达到较高SSIM，但本次结果仍弱于原90角LS+TV的像素精度；背景、前景和边缘诊断说明了这种指标差异。')
    bullet('正则化及多数扩展按GT NRMSE选参；本文结果对应所列迭代设置。4线程BM3D重跑观察到小幅数值变化，保存结果与核验记录均保留在仓库。')
    h('10.3 代码 参数与复现',2)
    add('link',label='项目首页与运行入口',url=github)
    add('link',label='环境 算法参数与复现命令',url=github+'/blob/main/docs/实验实施记录.md')
    add('link',label='完整实验结果与历史记录',url=github+'/blob/main/docs/基线实验结果.md')
    add('link',label='主表与汇总CSV',url=github+'/tree/main/outputs/final')
    page()
    h('附录 其他扩展结果')
    h('A.1 DIP与DIP加TV',2)
    dip=read('dip_aggregate.csv');rows=[]
    for name in ('SheppLogan64','XCAT128'):
        for k in (30,15):
            for snr in (20,30):
                for method in ('DIP','DIP_TV'):
                    r=next(r for r in dip if r['dataset']==name and int(r['angle_count'])==k and int(r['snr_label_db'])==snr and r['method']==method and r['checkpoint']=='fixed')
                    rows.append([SHORT[name],k,snr,method.replace('_','+'),
                        f"{float(r['nrmse_mean']):.4f}±{float(r['nrmse_std']):.4f}",
                        f"{float(r['psnr_db_mean']):.2f}±{float(r['psnr_db_std']):.2f}",
                        f"{float(r['ssim_mean']):.4f}±{float(r['ssim_std']):.4f}"])
    table('表A1 DIP固定2000步的三种子均值与样本标准差',['数据集','角度','SNR','方案','NRMSE↓','PSNR↑ dB','SSIM↑'],rows,[2.0,1.0,1.0,2.4,3.2,3.2,3.2])
    bullet('DIP与DIP+TV共48次拟合。DIP+TV在全部八个条件下的平均NRMSE均低于DIP；30角／30 dB的SL64结果由0.1815降至0.1029。')
    bullet('另完成24次留出投影DIP+TV拟合：8个条件中6个的验证截点平均NRMSE优于同次训练固定2000步，1个优于历史全投影训练结果。')
    h('A.2 TGV与小波先验',2)
    bullet('同求解器、同单个候选迭代预算下，非负TV在8个条件均优于无约束TV；TGV在5个条件的NRMSE低于非负TV。TGV与TV搜索分别为27组与9组。')
    bullet('小波＋TV在主20个条件中的11个比非负TV具有更低固定终点NRMSE，PnP-LC为14个。')
    page()
    h('参考文献')
    bib=json.loads((BUILD/'report_bibliography.json').read_text(encoding='utf-8'))
    for row in bib:add('reference',text=f"[{row['number']}] {row['text']}")


def font(run,size=11,bold=False,east='宋体',latin='Times New Roman'):
    run.font.name=latin;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for key,value in (('eastAsia',east),('ascii',latin),('hAnsi',latin),('cs',latin)):rf.set(qn('w:'+key),value)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rf.attrib.pop(qn('w:'+key),None)


def math_run(text):
    node=OxmlElement('m:r')
    mathpr=OxmlElement('m:rPr');normal=OxmlElement('m:nor');mathpr.append(normal);node.append(mathpr)
    pr=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts')
    for key,value in (('eastAsia','宋体'),('ascii','Times New Roman'),('hAnsi','Times New Roman')):rf.set(qn('w:'+key),value)
    pr.append(rf);node.append(pr)
    t=OxmlElement('m:t');t.text=text;node.append(t);return node


def append_math(parent,text):
    # 下标使用Word原生结构，避免把源稿中的x_gt等标记直接印进公式。
    pattern=r'([A-Za-z]+)_([A-Za-z0-9σ]+(?:,ᵢ)?)'
    cursor=0
    for match in re.finditer(pattern,text):
        if match.start()>cursor:parent.append(math_run(text[cursor:match.start()]))
        sub=OxmlElement('m:sSub');base=OxmlElement('m:e');index=OxmlElement('m:sub')
        base.append(math_run(match.group(1)));index.append(math_run(match.group(2).replace('ᵢ','i')))
        sub.extend([base,index]);parent.append(sub);cursor=match.end()
    if cursor<len(text):parent.append(math_run(text[cursor:]))


def ref_runs(para,text):
    for part in re.split(r'(https?://\S+)',text):
        if part.startswith(('http://','https://')):
            rel=para.part.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'18');pr.append(size);r.append(pr)
            rf=OxmlElement('w:rFonts')
            for key,value in (('eastAsia','宋体'),('ascii','Times New Roman'),('hAnsi','Times New Roman')):rf.set(qn('w:'+key),value)
            pr.append(rf)
            t=OxmlElement('w:t');t.text=part;r.append(t);link.append(r);para._p.append(link)
        else:font(para.add_run(part),9.3)


def math_para(doc,parts):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(6)
    om=OxmlElement('m:oMathPara');o=OxmlElement('m:oMath')
    for part in parts:
        if isinstance(part,str):append_math(o,part);continue
        if part[0]=='text':append_math(o,part[1])
        elif part[0]=='frac':
            f=OxmlElement('m:f');num=OxmlElement('m:num');den=OxmlElement('m:den')
            append_math(num,part[1]);append_math(den,part[2]);f.extend([num,den]);o.append(f)
    om.append(o);para._p.append(om)


def write_docx():
    doc=Document();sec=doc.sections[0]
    # 清除模板继承的段落装饰线；课程报告标题使用纯黑文字。
    for style in doc.styles:
        for border in list(style._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        rf=style._element.get_or_add_rPr().get_or_add_rFonts()
        for key,value in (('eastAsia','宋体'),('ascii','Times New Roman'),('hAnsi','Times New Roman'),('cs','Times New Roman')):rf.set(qn('w:'+key),value)
        for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rf.attrib.pop(qn('w:'+key),None)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.35);sec.right_margin=Cm(2.35)
    sec.header_distance=Cm(.7);sec.footer_distance=Cm(.9)
    for name in ('Normal','Title','Heading 1','Heading 2','Caption'):
        st=doc.styles[name];st.font.color.rgb=RGBColor(0,0,0);st.font.name='Times New Roman'
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.3;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for name,size in (('Title',20),('Heading 1',15),('Heading 2',12)):
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.styles['Caption'].font.size=Pt(9.3)
    doc.core_properties.title=TITLE;doc.core_properties.subject='Project 1 CT reconstruction coursework'
    doc.core_properties.author='';doc.core_properties.keywords='CT;LSQR;ART;TV;PnP;SSIM'
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for block in BLOCKS:
        kind=block['kind']
        if kind=='page':doc.add_page_break()
        elif kind=='h':
            level=block['level'];para=doc.add_paragraph(style='Title' if level==0 else f'Heading {level}')
            para.add_run(block['text']);para.paragraph_format.keep_with_next=True
            if level==0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif kind=='p':
            para=doc.add_paragraph();font(para.add_run(block['text']),10.5)
        elif kind=='bullet':
            para=doc.add_paragraph()
            para.paragraph_format.left_indent=Cm(.4);para.paragraph_format.first_line_indent=Cm(-.4)
            font(para.add_run('•  '+block['text']),10.5)
        elif kind=='toc':
            para=doc.add_paragraph();font(para.add_run('目录'),14,bold=True)
            doc.add_paragraph('[[TOC]]')
        elif kind=='link':
            para=doc.add_paragraph();font(para.add_run(block['label']+'：'),10.5)
            ref_runs(para,block['url'])
        elif kind=='reference':
            para=doc.add_paragraph();para.paragraph_format.line_spacing=1.12
            para.paragraph_format.space_after=Pt(5);ref_runs(para,block['text'])
        elif kind=='eq':math_para(doc,block['parts'])
        elif kind=='code':
            for line in block['text'].splitlines():
                para=doc.add_paragraph();para.paragraph_format.line_spacing=1.0
                font(para.add_run(line),8.4)
        elif kind=='fig':
            para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
            para.add_run().add_picture(str(DEST/block['path']),width=Cm(block['width']))
            cap=doc.add_paragraph(block['caption'],style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif kind=='table':
            compact=block['caption'].startswith('表13 ')
            cap=doc.add_paragraph(block['caption'],style='Caption');cap.paragraph_format.keep_with_next=True
            t=doc.add_table(rows=1, cols=len(block['headers']));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            widths=block['widths'] or [16/len(block['headers'])]*len(block['headers'])
            for i,width in enumerate(widths):t.columns[i].width=Cm(width);t.rows[0].cells[i].width=Cm(width)
            for cell,text in zip(t.rows[0].cells,block['headers']):cell.text=str(text)
            for row in block['rows']:
                cells=t.add_row().cells
                for i,(cell,value) in enumerate(zip(cells,row)):cell.width=Cm(widths[i]);cell.text=str(value)
            pr=t._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                node=OxmlElement('w:'+side);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
            pr.append(borders)
            for ri,row in enumerate(t.rows):
                rowpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');rowpr.append(cant)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader');rowpr.append(repeat)
                for ci,cell in enumerate(row.cells):
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                    for side in ('top','bottom','left','right'):
                        m=OxmlElement('w:'+side);m.set(qn('w:w'),('30' if compact else '55') if side in ('top','bottom') else '70');m.set(qn('w:type'),'dxa');margins.append(m)
                    tcpr.append(margins)
                    if ri==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EEF4');tcpr.append(shade)
                    for para in cell.paragraphs:
                        para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.line_spacing=1.06
                        if compact:
                            para.paragraph_format.line_spacing=1.0
                            snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');para._p.get_or_add_pPr().append(snap)
                        para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(0)
                        for run in para.runs:font(run,9.0 if len(t.rows)<30 else 8.9,bold=ri==0)
            if not compact:doc.add_paragraph().paragraph_format.space_after=Pt(2)
    path=DEST/(FILE_STEM+'.docx');doc.save(path)
    return path


def write_markdown():
    lines=[]
    for b in BLOCKS:
        k=b['kind']
        if k=='h':lines += [('#' if b['level']==0 else '#'*(b['level']+1))+' '+b['text'],'']
        elif k in ('p','reference'):lines += [b['text'],'']
        elif k=='bullet':lines += ['- '+b['text'],'']
        elif k=='link':lines += ['['+b['label']+']('+b['url']+')','']
        elif k=='toc':
            lines += ['## 目录','']
            for head in BLOCKS:
                if head['kind']=='h' and head['level']==1:
                    anchor=head['text'].lower().replace(' ','-').replace('.','')
                    lines += ['- ['+head['text']+'](#'+anchor+')']
            lines += ['']
        elif k=='eq':lines += ['> '+b['text'],'']
        elif k=='code':lines += ['```powershell',b['text'],'```','']
        elif k=='fig':lines += ['!['+b['caption']+']('+b['path']+')','']
        elif k=='table':
            lines += [b['caption'],'','| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |']
            lines += ['| '+' | '.join(str(x) for x in row)+' |' for row in b['rows']];lines += ['']
    path=DEST/(FILE_STEM+'.md');path.write_text('\n'.join(lines),encoding='utf-8');return path


def main():
    DEST.mkdir(exist_ok=True);BUILD.mkdir(parents=True,exist_ok=True)
    compose();d=write_docx();m=write_markdown()
    (BUILD/'report_tables.json').write_text(json.dumps(TABLE_DATA,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'docx':str(d),'markdown':str(m),'tables':len(TABLE_DATA),'blocks':len(BLOCKS)},ensure_ascii=False))


if __name__=='__main__':main()
